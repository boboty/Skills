"""台账工作簿存储层: 加载/写入/备份/锁/元数据/审计/导入登记。

设计要点:
- 单一 Excel 文件即整本台账; 每张工作表首行固定表头(见 schema.COLUMNS)。
- 所有写入走 append/update + save(); save() 前自动备份, 原子替换, 写后校验。
- 并发: 写模式打开时创建 <path>.lock (进程级, 含 pid); PID 不存活视为陈旧锁可回收。
- 数值列写 Decimal(openpyxl 落 float), 读回一律 Decimal(str(v)) 还原, 保证精确。
- 业务日期/期间用 ISO 字符串, 避免 Excel 日期类型时区歧义。
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import os
import shutil
import time
from pathlib import Path

import openpyxl

from . import schema as S
from .decimal_utils import as_decimal, money, price, qty

class LedgerError(Exception):
    """业务/校验错误。code 对应 business-rules.md 下篇异常清单编号。"""

    def __init__(self, code: str, message: str):
        super().__init__(f"[{code}] {message}")
        self.code = code
        self.message = message


def now_str() -> str:
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def today_str() -> str:
    return _dt.date.today().isoformat()


def is_period_valid(period: str) -> bool:
    try:
        _dt.datetime.strptime(period, "%Y-%m")
        return True
    except ValueError:
        return False


def is_date_valid(d: str) -> bool:
    try:
        _dt.datetime.strptime(d, "%Y-%m-%d")
        return True
    except ValueError:
        return False


def period_of(date_str: str) -> str:
    return date_str[:7]


def prev_period(period: str) -> str:
    y, m = int(period[:4]), int(period[5:7])
    if m == 1:
        return f"{y - 1}-12"
    return f"{y}-{m - 1:02d}"


def next_period(period: str) -> str:
    y, m = int(period[:4]), int(period[5:7])
    if m == 12:
        return f"{y + 1}-01"
    return f"{y}-{m + 1:02d}"


def sha256_of_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


class _Batch:
    """with store.batch(): 内的所有 save() 挂起; 退出时一次性写盘。"""

    def __init__(self, st: "Store"):
        self.st = st

    def __enter__(self):
        self.st._defer_save = True
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.st._defer_save = False
        if exc_type is None:
            self.st.save()  # 唯一一次真实写盘(含备份/原子替换/写后校验)
        return False


class Store:
    """一本台账 = 一个 xlsx 文件 + 缓存的行数据。"""

    def __init__(self, path: str, wb: openpyxl.Workbook, lock_owned: bool = False):
        self.path = str(Path(path).resolve())
        self.wb = wb
        self.lock_owned = lock_owned
        self._defer_save = False

    # ------------------------------------------------------------ 打开/创建
    @classmethod
    def create(cls, path: str, company: str, start_period: str,
               operator: str = "system") -> "Store":
        if not is_period_valid(start_period):
            raise LedgerError("E08", f"起始期间格式应为 {S.PERIOD_FMT}: {start_period}")
        if os.path.exists(path):
            raise LedgerError("E16", f"文件已存在, 拒绝覆盖: {path}")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet in S.ALL_SHEETS:
            ws = wb.create_sheet(sheet)
            ws.append(S.headers(sheet))
        st = cls(path, wb)
        # 新建即视为独占该文件(允许 save; close 时释放锁文件, 不存在则忽略)
        st.lock_owned = True
        st.set_meta(S.META_LEDGER_ID, hashlib.sha1(str(time.time_ns()).encode()).hexdigest()[:12])
        st.set_meta(S.META_COMPANY, company)
        st.set_meta(S.META_CURRENCY, S.CURRENCY)
        st.set_meta(S.META_COST_METHOD, S.COST_METHOD)
        st.set_meta(S.META_PRICE_PREC, S.PREC_PRICE)
        st.set_meta(S.META_MONEY_PREC, S.PREC_MONEY)
        st.set_meta(S.META_QTY_PREC, S.PREC_QTY)
        st.set_meta(S.META_ROUND_MODE, S.ROUND_MODE)
        st.set_meta(S.META_START_PERIOD, start_period)
        st.set_meta(S.META_VERSION, "1.0.0")
        st.set_meta(S.META_CREATED_AT, now_str())
        st.set_meta(S.META_WRITES, "0")
        st.set_period(start_period, S.PERIOD_OPEN, operator, "初始化")
        st.audit("init", operator, f"初始化台账 {company} 起始期间 {start_period}",
                 "创建工作表与元数据", path, "成功")
        st.save()
        return st

    @classmethod
    def open(cls, path: str, for_write: bool = False) -> "Store":
        p = str(Path(path).resolve())
        if not os.path.exists(p):
            raise LedgerError("E16", f"台账文件不存在: {p}")
        lock_owned = False
        if for_write:
            lock_owned = cls._acquire_lock(p)
        try:
            wb = openpyxl.load_workbook(p, data_only=False)
        except Exception as e:  # noqa: BLE001
            if lock_owned:
                cls._release_lock(p)
            raise LedgerError("E10", f"Excel 读取失败: {e}") from e
        st = cls(p, wb, lock_owned=lock_owned)
        st._check_structure()
        return st

    @staticmethod
    def _lock_path(path: str) -> str:
        return path + ".lock"

    @classmethod
    def _acquire_lock(cls, path: str) -> bool:
        lp = cls._lock_path(path)
        if os.path.exists(lp):
            try:
                pid = int(open(lp).read().split()[0])
                os.kill(pid, 0)  # noqa: PLW1509 存活检测
                raise LedgerError("E17", f"台账正被其他进程写入(锁文件 {lp}), 不支持并发编辑")
            except ProcessLookupError:
                os.remove(lp)  # 陈旧锁回收
            except ValueError:
                os.remove(lp)
        with open(lp, "w") as f:
            f.write(f"{os.getpid()} {now_str()}\n")
        return True

    @classmethod
    def _release_lock(cls, path: str) -> None:
        lp = cls._lock_path(path)
        try:
            if os.path.exists(lp):
                os.remove(lp)
        except OSError:
            pass

    def _check_structure(self) -> None:
        have = {ws.title for ws in self.wb.worksheets}
        missing = [s for s in S.ALL_SHEETS if s not in have]
        if missing:
            raise LedgerError("E16", f"台账缺少工作表: {missing}, 文件可能不是有效台账")

    # ------------------------------------------------------------ 行读写
    def read_all(self, sheet: str) -> list[dict]:
        """读整表为 dict 列表(值经 Decimal 还原)。"""
        ws = self.wb[sheet]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return []
        out = []
        for i, r in enumerate(rows, start=2):  # 数据行号从 2 开始(表头占 1)
            if all(c is None for c in r):
                continue
            d = {}
            for idx, h in enumerate(header):
                if idx >= len(r):
                    d[h] = None
                    continue
                v = r[idx]
                if v is not None and h in S.DECIMAL_COLS:
                    v = as_decimal(v)
                d[h] = v
            d["_row"] = i  # 物理行号, 用于原位更新
            out.append(d)
        return out

    def append_row(self, sheet: str, values: dict) -> None:
        ws = self.wb[sheet]
        headers = S.headers(sheet)
        row = [values.get(h, "") for h in headers]
        ws.append(row)

    def update_cell(self, sheet: str, data_row: int, header: str, value) -> None:
        ws = self.wb[sheet]
        headers = S.headers(sheet)
        col = headers.index(header) + 1
        ws.cell(row=data_row, column=col, value=value)

    def next_id(self, sheet: str, prefix: str) -> str:
        id_field = self._id_field(sheet)
        max_n = 0
        for d in self.read_all(sheet):
            rid = str(d.get(id_field) or "")
            if rid.startswith(prefix):
                try:
                    max_n = max(max_n, int(rid[len(prefix):]))
                except ValueError:
                    pass
        return f"{prefix}{max_n + 1:06d}"

    @staticmethod
    def _id_field(sheet: str) -> str:
        return {
            S.SHEET_OPENING: "期初ID",
            S.SHEET_IN: "入库ID",
            S.SHEET_OUT: "出库ID",
            S.SHEET_TRANSFER: "调拨ID",
            S.SHEET_ADJUST: "调整ID",
            S.SHEET_FEE: "费用ID",
            S.SHEET_AUDIT: "操作ID",
            S.SHEET_IMPORTS: "导入ID",
        }.get(sheet, "")

    # ------------------------------------------------------------ 元数据
    def meta(self, key: str, default=None):
        for d in self.read_all(S.SHEET_META):
            if d.get("键") == key:
                return d.get("值", default)
        return default

    def set_meta(self, key: str, value) -> None:
        for d in self.read_all(S.SHEET_META):
            if d.get("键") == key:
                self.update_cell(S.SHEET_META, d["_row"], "值", value)
                return
        self.append_row(S.SHEET_META, {"键": key, "值": value, "备注": ""})

    # ------------------------------------------------------------ 期间状态
    def get_period(self, period: str) -> dict | None:
        for d in self.read_all(S.SHEET_PERIODS):
            if d.get("期间") == period:
                return d
        return None

    def set_period(self, period: str, status: str, operator: str,
                   remark: str = "") -> None:
        d = self.get_period(period)
        ver = int(d["核算版本"]) if d and d.get("核算版本") else 0
        if status == S.PERIOD_CLOSED:
            ver += 1
        if d:
            self.update_cell(S.SHEET_PERIODS, d["_row"], "状态", status)
            self.update_cell(S.SHEET_PERIODS, d["_row"], "核算版本", ver)
            self.update_cell(S.SHEET_PERIODS, d["_row"], "更新时间", now_str())
            self.update_cell(S.SHEET_PERIODS, d["_row"], "操作人", operator)
            if remark:
                self.update_cell(S.SHEET_PERIODS, d["_row"], "备注", remark)
        else:
            self.append_row(S.SHEET_PERIODS, {
                "期间": period, "状态": status, "核算版本": ver,
                "更新时间": now_str(), "操作人": operator, "备注": remark,
            })

    def ensure_period_open(self, period: str) -> None:
        if not is_period_valid(period):
            raise LedgerError("E08", f"期间格式应为 {S.PERIOD_FMT}: {period}")
        d = self.get_period(period)
        if d is None:
            raise LedgerError("E08", f"期间 {period} 未登记(超出台账期间范围)")
        if d.get("状态") == S.PERIOD_CLOSED:
            raise LedgerError("E15", f"期间 {period} 已结账, 禁止直接修改; 如需修改请走受控重开")

    def latest_period(self) -> str:
        periods = [d["期间"] for d in self.read_all(S.SHEET_PERIODS) if d.get("期间")]
        return max(periods) if periods else self.meta(S.META_START_PERIOD)

    # ------------------------------------------------------------ 审计/导入登记
    def audit(self, op_type: str, operator: str, summary: str,
              detail: str = "", ref_key: str = "", result: str = "成功") -> str:
        op_id = self.next_id(S.SHEET_AUDIT, "OP-")
        self.append_row(S.SHEET_AUDIT, {
            "操作ID": op_id, "时间": now_str(), "操作类型": op_type,
            "操作人": operator, "指令摘要": summary, "变更摘要": detail,
            "关联键": ref_key, "结果": result,
        })
        return op_id

    def register_import(self, file_name: str, file_hash: str, file_size: int,
                        imp_type: str, ok: int, reject: int, pending: int,
                        operator: str) -> str:
        imp_id = self.next_id(S.SHEET_IMPORTS, "IMP-")
        batch = f"B{time.strftime('%Y%m%d%H%M%S')}-{imp_id.split('-')[1]}"
        self.append_row(S.SHEET_IMPORTS, {
            "导入ID": imp_id, "批次号": batch, "文件名": file_name,
            "文件摘要": file_hash, "文件大小": file_size, "导入类型": imp_type,
            "成功条数": ok, "拒绝条数": reject, "待确认条数": pending,
            "导入时间": now_str(), "操作人": operator, "状态": S.IMPORT_OK,
        })
        return batch

    def find_import(self, file_hash: str, imp_type: str) -> dict | None:
        for d in self.read_all(S.SHEET_IMPORTS):
            if d.get("文件摘要") == file_hash and d.get("导入类型") == imp_type:
                return d
        return None

    # ------------------------------------------------------------ 保存
    def save(self) -> None:
        if self._defer_save:
            return  # batch 挂起: 由 batch 退出时一次性写盘(原子提交)
        if not self.lock_owned:
            raise LedgerError("E17", "只读打开, 禁止保存")
        writes = int(self.meta(S.META_WRITES, "0"))
        self.set_meta(S.META_WRITES, str(writes + 1))
        # 1) 备份当前正式文件
        if os.path.exists(self.path):
            bak = Path(self.path).parent / "backups" / (
                f"{Path(self.path).stem}_{time.strftime('%Y%m%d_%H%M%S')}_"
                f"{sha256_of_file(self.path)[:8]}.xlsx")
            bak.parent.mkdir(exist_ok=True)
            shutil.copy2(self.path, bak)
        # 2) 原子替换写入
        tmp = self.path + f".tmp{os.getpid()}"
        try:
            self.wb.save(tmp)
            os.replace(tmp, self.path)
        except Exception as e:  # noqa: BLE001
            if os.path.exists(tmp):
                os.remove(tmp)
            raise LedgerError("E10", f"Excel 写入失败: {e}") from e
        # 3) 写后校验: 重载并确认写入次数
        v = openpyxl.load_workbook(self.path, data_only=True)
        found = False
        for row in v[S.SHEET_META].iter_rows(values_only=True):
            if row and row[0] == S.META_WRITES and row[1] is not None and str(row[1]) == str(writes + 1):
                found = True
                break
        if not found:
            raise LedgerError("E10", "写入后校验失败: 元数据未更新, 文件可能损坏")

    def batch(self) -> "_Batch":
        """原子批处理: 挂起 save() 直至退出; 正常退出写盘一次, 异常退出不写盘。"""
        if self._defer_save:
            raise LedgerError("E10", "不允许嵌套 batch(原子提交)")
        return _Batch(self)

    def close(self) -> None:
        if self.lock_owned:
            self._release_lock(self.path)
            self.lock_owned = False

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.close()